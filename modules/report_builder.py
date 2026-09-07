"""
report_builder.py
-------------------
Turns the final report text (either LLM-written or the deterministic
fallback built here) plus the structured difference list into downloadable
TXT, DOCX, and PDF files.

The deterministic fallback (`build_fallback_report_text`) exists so the
tool remains fully usable with zero LLM API key configured -- it produces
the same section structure, just with template sentences instead of
free-form AI prose.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List

from .difference_builder import DifferenceRecord, summarize
from .utils import get_logger

logger = get_logger(__name__)

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    _DOCX_AVAILABLE = True
except Exception:  # pragma: no cover
    _DOCX_AVAILABLE = False
    logger.warning("python-docx not installed -- DOCX report export disabled.")

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    _REPORTLAB_AVAILABLE = True
except Exception:  # pragma: no cover
    _REPORTLAB_AVAILABLE = False
    logger.warning("reportlab not installed -- PDF report export disabled.")


SEVERITY_COLOR_HEX = {"high": "C0392B", "medium": "D68910", "low": "2E86C1", "none": "27AE60"}

# Plain-language translation of the internal status codes, for readers who
# aren't DTP specialists (e.g. a PM or client skimming the report).
STATUS_PLAIN_TEXT = {
    "MATCH": "No issue -- matches the source.",
    "MOVED": "This element shifted position compared to the source.",
    "PAGE_MOVED": "This element ended up on a different page than in the source.",
    "RESIZED": "This element's size changed compared to the source.",
    "OVERFLOW": "The text no longer fits its box and is being cut off or spilling out.",
    "UNDERFLOW": "This element is taking up noticeably less space than expected -- possible missing content.",
    "ALIGNMENT_CHANGED": "This element's alignment (e.g. left edge) no longer lines up with the source.",
    "MARGIN_CHANGED": "The page margins, header, or footer don't match the source.",
    "CROPPED": "This image appears to be cropped compared to the source.",
    "ROTATED": "This element's rotation differs from the source.",
    "SPLIT": "A table that was one piece in the source is now split across multiple pieces.",
    "MERGED": "Multiple separate tables in the source now appear merged into one.",
    "MISSING": "This element from the source is missing in the translated file.",
    "EXTRA": "This element appears in the translated file but has no match in the source.",
    "COLOR_CHANGED": "This text's color no longer matches the source.",
    "FORMATTING_CHANGED": "This text's bold/italic styling no longer matches the source.",
    "BULLET_MISMATCH": "The number of bullet points/list items here doesn't match the source -- one may be missing or an extra one may have been added.",
    "STRUCTURE_CHANGED": "This table's number of rows and/or columns doesn't match the source.",
}

OBJECT_TYPE_PLAIN_TEXT = {
    "paragraph": "Text block",
    "image": "Image",
    "table": "Table",
    "page": "Page-level",
}

SEVERITY_PLAIN_TEXT = {
    "high": "Must fix before publishing",
    "medium": "Should review",
    "low": "Minor, low priority",
    "none": "No action needed",
}


def _plain_issue_line(r: "DifferenceRecord") -> str:
    """One human-readable sentence for a single issue, no jargon codes required to understand it."""
    what = OBJECT_TYPE_PLAIN_TEXT.get(r.object_type, r.object_type.capitalize())
    what_happened = STATUS_PLAIN_TEXT.get(r.status, r.status.replace("_", " ").title())
    priority = SEVERITY_PLAIN_TEXT.get(r.severity, r.severity)
    return f"- **{what}** -- {what_happened} ({priority}). {r.recommendation}"


# --------------------------------------------------------------------------
# Deterministic fallback text (used if no LLM key is configured)
# --------------------------------------------------------------------------

def build_fallback_report_text(
    records: List[DifferenceRecord],
    source_name: str,
    target_name: str,
    source_lang: str,
    target_lang: str,
    page_count: int,
) -> str:
    summary = summarize(records)
    lines = []
    lines.append("## Document Information")
    lines.append(f"- Source file: {source_name}")
    lines.append(f"- Target file: {target_name}")
    lines.append(f"- Source language: {source_lang}")
    lines.append(f"- Target language: {target_lang}")
    lines.append(f"- Page count: {page_count}")
    lines.append(f"- Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append("")
    lines.append("## Overall Status")
    status_plain = {
        "PASS": "🟢 PASS -- the layout matches the source. No changes needed before publishing.",
        "REVIEW REQUIRED": "🟡 REVIEW REQUIRED -- some layout differences were found. A designer should look these over before publishing.",
        "FAIL": "🔴 FAIL -- one or more serious layout problems were found. This file should not be published until they're fixed.",
    }
    lines.append(status_plain.get(summary["overall_status"], f"**{summary['overall_status']}**"))
    lines.append("")
    lines.append("## QA Summary")
    lines.append(f"- Total issues found: {summary['total_issues']}")
    lines.append(f"- Must fix before publishing: {summary['by_severity'].get('high', 0)}")
    lines.append(f"- Should review: {summary['by_severity'].get('medium', 0)}")
    lines.append(f"- Minor / low priority: {summary['by_severity'].get('low', 0)}")
    lines.append("")
    lines.append("## Page-wise Issues")
    if not records:
        lines.append("No layout issues detected -- every page matches the source.")
    else:
        pages = sorted(set(r.page for r in records))
        for page in pages:
            lines.append(f"### Page {page}")
            for r in [x for x in records if x.page == page]:
                lines.append(_plain_issue_line(r))
    lines.append("")
    lines.append("## Designer Action Items")
    high = [r for r in records if r.severity == "high"]
    if high:
        for r in high:
            lines.append(f"- [Page {r.page}] {r.recommendation}")
    else:
        lines.append("- No critical (must-fix) action items.")
    lines.append("")
    lines.append("## Final Remark")
    if summary["overall_status"] == "PASS":
        lines.append("Layout is consistent with the source document. Ready for publishing pending final visual check.")
    elif summary["overall_status"] == "REVIEW REQUIRED":
        lines.append("Minor to moderate layout deviations were found. Designer review recommended before publishing.")
    else:
        lines.append("Critical layout issues were found. This file should not be published until resolved.")
    return "\n".join(lines)


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------

def export_txt(report_text: str) -> bytes:
    return report_text.encode("utf-8")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def export_docx(report_text: str, records: List[DifferenceRecord]) -> bytes:
    if not _DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed. Run `pip install python-docx`.")

    doc = Document()
    title = doc.add_heading("AI DTP QA REPORT", level=0)

    for raw_line in report_text.split("\n"):
        line = raw_line.rstrip()
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip().startswith("**") and line.strip().endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip().strip("*"))
            run.bold = True
        elif line.strip():
            doc.add_paragraph(line)

    if records:
        doc.add_heading("Detailed Issue Table", level=1)
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Page", "Type", "Status", "Severity", "Confidence", "Recommendation"]):
            hdr[i].text = h
        for r in records:
            row = table.add_row().cells
            row[0].text = str(r.page)
            row[1].text = r.object_type
            row[2].text = r.status
            row[3].text = r.severity
            row[4].text = f"{r.confidence:.2f}"
            row[5].text = r.recommendation

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def export_pdf(report_text: str, records: List[DifferenceRecord]) -> bytes:
    if not _REPORTLAB_AVAILABLE:
        raise RuntimeError("reportlab is not installed. Run `pip install reportlab`.")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], spaceAfter=8)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], spaceAfter=6)
    body = styles["BodyText"]

    story = [Paragraph("AI DTP QA REPORT", styles["Title"]), Spacer(1, 12)]

    for raw_line in report_text.split("\n"):
        line = raw_line.rstrip()
        if line.startswith("## "):
            story.append(Paragraph(line[3:], h1))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], h2))
        elif line.startswith("- "):
            story.append(Paragraph(f"&bull; {line[2:]}", body))
        elif line.strip():
            story.append(Paragraph(line.replace("**", ""), body))
        else:
            story.append(Spacer(1, 6))

    if records:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Detailed Issue Table", h1))
        data = [["Page", "Type", "Status", "Severity", "Conf.", "Recommendation"]]
        for r in records:
            data.append([str(r.page), r.object_type, r.status, r.severity, f"{r.confidence:.2f}",
                         Paragraph(r.recommendation, body)])
        tbl = Table(data, colWidths=[1.4 * cm, 2.0 * cm, 2.6 * cm, 1.8 * cm, 1.6 * cm, 7.0 * cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(tbl)

    doc.build(story)
    return buf.getvalue()


def export_all(
    report_text: str,
    records: List[DifferenceRecord],
) -> Dict[str, bytes]:
    """Convenience helper returning whichever formats are available in this environment."""
    outputs: Dict[str, bytes] = {"txt": export_txt(report_text)}
    if _DOCX_AVAILABLE:
        outputs["docx"] = export_docx(report_text, records)
    if _REPORTLAB_AVAILABLE:
        outputs["pdf"] = export_pdf(report_text, records)
    return outputs
